from __future__ import annotations

import difflib
import io
import json
import os
import re
import sqlite3
import statistics
import time
import uuid
from dataclasses import dataclass, asdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
import streamlit as st
from docx import Document

try:
    from pangram import Pangram
except Exception:
    Pangram = None


APP_TITLE = "Pangram Experiment Lab"
APP_VERSION = "v2.0 · Streamlit Cloud"
MIN_PANGRAM_WORDS = 50
DB_PATH = Path(__file__).with_name("pangram_microscope.db")
DEFAULT_SAMPLE_SIZES = [150]
ALL_SAMPLE_SIZES = [50, 75, 100, 125, 150, 200, 250, 300, 400, 500, 750, 1000]
CALIBRATED_WINDOW_WORDS = 150
DEFAULT_OVERLAP_PCT = 0
QUICK_DEFAULT_MAX_WINDOWS = 20
CALIBRATION_DEFAULT_MAX_WINDOWS = 4
PANGRAM_REALTIME_RATE_PER_100_WORDS = 0.05
PANGRAM_BULK_DISCOUNT = 0.20
PANGRAM_BULK_RATE_PER_100_WORDS = PANGRAM_REALTIME_RATE_PER_100_WORDS * (1.0 - PANGRAM_BULK_DISCOUNT)
COST_WARNING_THRESHOLD = 5.00
EXPERIMENT_WINDOW_WORDS = 500
EXPERIMENT_OVERLAP_PCT = 50
EXPERIMENT_MAX_WINDOWS_PER_FILE = 20
DEFAULT_STRUCTURE_SIMILARITY_LIMIT = 0.82


@dataclass
class SourceDoc:
    name: str
    expected_label: str
    text: str


@dataclass
class TextWindow:
    window_id: str
    source_name: str
    expected_label: str
    target_words: int
    actual_words: int
    sentence_start: int
    sentence_end: int
    text: str


# -------------------------
# Text handling
# -------------------------

WORD_RE = re.compile(r"\S+")


def count_words(text: str) -> int:
    return len(WORD_RE.findall(text or ""))


def clean_text(text: str) -> str:
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    # Keep paragraph boundaries, but normalize internal whitespace.
    paragraphs = []
    for p in re.split(r"\n\s*\n", text):
        p = re.sub(r"[ \t\f\v]+", " ", p).strip()
        if p:
            paragraphs.append(p)
    return "\n\n".join(paragraphs)


def extract_text_from_upload(uploaded_file) -> str:
    suffix = Path(uploaded_file.name).suffix.lower()
    raw = uploaded_file.getvalue()

    if suffix == ".docx":
        doc = Document(io.BytesIO(raw))
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        return clean_text("\n\n".join(parts))

    if suffix in {".txt", ".md"}:
        for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                return clean_text(raw.decode(encoding))
            except UnicodeDecodeError:
                continue
        return clean_text(raw.decode("utf-8", errors="replace"))

    raise ValueError(f"Unsupported file type: {suffix}. Use DOCX, TXT, or MD.")


def split_sentences(text: str) -> list[str]:
    """A lightweight fiction-friendly sentence splitter.

    Pangram 4 is designed for complete-sentence prose. We therefore build test
    windows on sentence boundaries rather than cutting at an exact word index.
    This is deliberately dependency-light; it is not intended as a linguistic
    parser.
    """
    text = clean_text(text)
    if not text:
        return []

    sentences: list[str] = []
    # Work paragraph by paragraph so paragraph breaks remain natural stopping points.
    for paragraph in re.split(r"\n\s*\n", text):
        paragraph = paragraph.strip()
        if not paragraph:
            continue

        # Capture through sentence-ending punctuation plus closing quote/bracket.
        # If a paragraph has no terminal punctuation, keep it as one unit.
        matches = re.findall(
            r".+?(?:[.!?]+(?:[\"'”’)]*)?(?=\s+|$)|$)",
            paragraph,
            flags=re.S,
        )
        for item in matches:
            item = item.strip()
            if item:
                sentences.append(item)

    return sentences


def _advance_start(sent_word_counts: list[int], start: int, stride_words: int) -> int:
    if start >= len(sent_word_counts) - 1:
        return len(sent_word_counts)
    total = 0
    i = start
    while i < len(sent_word_counts) and total < stride_words:
        total += sent_word_counts[i]
        i += 1
    return max(start + 1, i)


def build_sentence_windows(
    doc: SourceDoc,
    target_words: int,
    overlap_fraction: float,
) -> list[TextWindow]:
    sentences = split_sentences(doc.text)
    if not sentences:
        return []

    counts = [count_words(s) for s in sentences]
    stride_words = max(1, int(round(target_words * (1.0 - overlap_fraction))))
    min_acceptable = MIN_PANGRAM_WORDS

    windows: list[TextWindow] = []
    start = 0
    ordinal = 0

    while start < len(sentences):
        total = 0
        end = start
        while end < len(sentences) and total < target_words:
            total += counts[end]
            end += 1

        if total < min_acceptable:
            break

        ordinal += 1
        text = " ".join(sentences[start:end]).strip()
        safe_name = re.sub(r"[^A-Za-z0-9_-]+", "_", Path(doc.name).stem)[:45]
        wid = f"{safe_name}__{target_words}w__{ordinal:04d}"
        windows.append(
            TextWindow(
                window_id=wid,
                source_name=doc.name,
                expected_label=doc.expected_label,
                target_words=target_words,
                actual_words=count_words(text),
                sentence_start=start + 1,
                sentence_end=end,
                text=text,
            )
        )

        start = _advance_start(counts, start, stride_words)

    return windows


def evenly_cap(items: list[TextWindow], cap: int) -> list[TextWindow]:
    if cap <= 0 or len(items) <= cap:
        return items
    if cap == 1:
        return [items[len(items) // 2]]
    indexes = [round(i * (len(items) - 1) / (cap - 1)) for i in range(cap)]
    seen = set()
    selected = []
    for idx in indexes:
        if idx not in seen:
            selected.append(items[idx])
            seen.add(idx)
    return selected


def make_calibration_windows(
    docs: list[SourceDoc],
    sample_sizes: list[int],
    overlap_fraction: float,
    cap_per_source_size: int,
) -> list[TextWindow]:
    out: list[TextWindow] = []
    for doc in docs:
        for size in sorted(sample_sizes):
            candidates = build_sentence_windows(doc, size, overlap_fraction)
            out.extend(evenly_cap(candidates, cap_per_source_size))
    return out


# -------------------------
# Cost estimation
# -------------------------


def estimate_bulk_cost(windows: list[TextWindow]) -> dict[str, float | int]:
    """Estimate Pangram bulk cost using started 100-word blocks per request.

    Assumption: $0.05 per started 100-word block less a 20% bulk discount,
    for an effective $0.04 per started 100-word block.
    """
    requests = len(windows)
    actual_words = sum(w.actual_words for w in windows)
    billing_blocks = sum(max(1, (w.actual_words + 99) // 100) for w in windows)
    billable_words = billing_blocks * 100
    estimated_cost = billing_blocks * PANGRAM_BULK_RATE_PER_100_WORDS
    return {
        "requests": requests,
        "actual_words": actual_words,
        "billing_blocks": billing_blocks,
        "billable_words": billable_words,
        "estimated_cost": estimated_cost,
    }


def show_cost_estimate(windows: list[TextWindow], *, key: str) -> None:
    if not windows:
        return
    est = estimate_bulk_cost(windows)
    st.info(
        f"Estimated Pangram bulk cost: **${est['estimated_cost']:.2f}** · "
        f"{est['requests']:,} requests · {est['actual_words']:,} actual words · "
        f"{est['billable_words']:,} billable words."
    )
    st.caption(
        "Estimate assumes $0.05 per started 100-word block with a 20% bulk discount "
        "($0.04 per block). Pangram bills each request separately, so sentence-aligned windows can round up."
    )
    if est["estimated_cost"] >= COST_WARNING_THRESHOLD:
        st.warning(
            f"Cost guardrail: this run is estimated at ${est['estimated_cost']:.2f}, "
            f"which is at or above the ${COST_WARNING_THRESHOLD:.2f} warning threshold."
        )


# -------------------------
# Pangram API
# -------------------------


def get_secret_api_key() -> str:
    env_key = os.getenv("PANGRAM_API_KEY", "").strip()
    if env_key:
        return env_key
    try:
        return str(st.secrets.get("PANGRAM_API_KEY", "")).strip()
    except Exception:
        return ""


def connect_pangram(api_key: str) -> tuple[Any, list[str]]:
    if Pangram is None:
        raise RuntimeError(
            "The pangram-sdk package is not installed. Make sure requirements.txt is in the GitHub repo, then reboot the Streamlit app."
        )
    client = Pangram(api_key=api_key) if api_key else Pangram()
    models = client.list_models()
    return client, list(models)


def weighted_window_metric(result: dict[str, Any], field: str) -> float | None:
    windows = result.get("windows") or []
    vals = []
    weights = []
    for w in windows:
        value = w.get(field)
        if value is None:
            continue
        weight = w.get("word_count") or count_words(w.get("text", "")) or 1
        vals.append(float(value))
        weights.append(float(weight))
    if not vals:
        return None
    return sum(v * wt for v, wt in zip(vals, weights)) / sum(weights)


def summarize_result(result: dict[str, Any]) -> dict[str, Any]:
    windows = result.get("windows") or []
    confidences = [str(w.get("confidence", "")) for w in windows if w.get("confidence")]
    humanizer_scores = [
        float(w["humanizer_score"])
        for w in windows
        if w.get("humanizer_score") is not None
    ]
    return {
        "prediction": result.get("prediction_short"),
        "headline": result.get("headline"),
        "fraction_ai": result.get("fraction_ai"),
        "fraction_ai_assisted": result.get("fraction_ai_assisted"),
        "fraction_human": result.get("fraction_human"),
        "num_ai_segments": result.get("num_ai_segments"),
        "num_ai_assisted_segments": result.get("num_ai_assisted_segments"),
        "num_human_segments": result.get("num_human_segments"),
        "mean_ai_involvement": weighted_window_metric(result, "ai_assistance_score"),
        "max_humanizer_score": max(humanizer_scores) if humanizer_scores else None,
        "any_humanized": any(bool(w.get("is_humanized")) for w in windows),
        "window_confidence": ", ".join(sorted(set(confidences))),
        "version": result.get("version"),
        "dashboard_link": result.get("dashboard_link"),
    }


def run_bulk_scan(
    client: Any,
    model: str,
    windows: list[TextWindow],
    batch_size: int = 200,
    timeout: float = 3600,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Submit one or more Pangram bulk jobs and return successful and failed rows."""
    successes: list[dict[str, Any]] = []
    failures: list[dict[str, Any]] = []

    meta_by_id = {w.window_id: w for w in windows}
    progress = st.progress(0, text="Submitting Pangram bulk job…")
    total_batches = max(1, (len(windows) + batch_size - 1) // batch_size)

    for batch_no, offset in enumerate(range(0, len(windows), batch_size), start=1):
        chunk = windows[offset : offset + batch_size]
        items = [{"id": w.window_id, "text": w.text} for w in chunk]

        bulk = client.submit_bulk(items=items, model=model)
        bulk_id = bulk["bulk_id"]
        progress.progress(
            min(0.85, (batch_no - 1) / total_batches + 0.05),
            text=f"Pangram batch {batch_no}/{total_batches}: waiting for results…",
        )
        status = client.wait_for_bulk(bulk_id, timeout=timeout, poll_interval=0.5)
        results = client.get_bulk_results(bulk_id)

        for item in results.get("items", []):
            item_id = item.get("id")
            result = item.get("result")
            if result is None:
                failures.append(
                    {
                        "window_id": item_id,
                        "error": item.get("error") or f"No result; stage={item.get('stage')}",
                        "bulk_id": bulk_id,
                    }
                )
                continue

            meta = meta_by_id.get(item_id)
            if meta is None:
                failures.append(
                    {"window_id": item_id, "error": "Unknown result ID", "bulk_id": bulk_id}
                )
                continue

            row = asdict(meta)
            row.update(summarize_result(result))
            row["bulk_id"] = bulk_id
            row["raw_result"] = result
            successes.append(row)

        for failed in results.get("failed_items", []):
            failures.append(
                {
                    "window_id": failed.get("id"),
                    "error": failed.get("error") or "Pangram bulk item failed",
                    "bulk_id": bulk_id,
                }
            )

        progress.progress(
            min(0.98, batch_no / total_batches),
            text=f"Pangram batch {batch_no}/{total_batches} complete.",
        )

    progress.progress(1.0, text="Pangram analysis complete.")
    time.sleep(0.15)
    progress.empty()
    return successes, failures


# -------------------------
# Persistence
# -------------------------


def init_db() -> None:
    with sqlite3.connect(DB_PATH) as con:
        con.execute(
            """
            CREATE TABLE IF NOT EXISTS scan_results (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                experiment_id TEXT NOT NULL,
                experiment_name TEXT,
                run_at TEXT NOT NULL,
                mode TEXT NOT NULL,
                model TEXT,
                source_name TEXT,
                expected_label TEXT,
                window_id TEXT,
                target_words INTEGER,
                actual_words INTEGER,
                sentence_start INTEGER,
                sentence_end INTEGER,
                prediction TEXT,
                headline TEXT,
                fraction_ai REAL,
                fraction_ai_assisted REAL,
                fraction_human REAL,
                mean_ai_involvement REAL,
                max_humanizer_score REAL,
                any_humanized INTEGER,
                window_confidence TEXT,
                version TEXT,
                dashboard_link TEXT,
                text TEXT,
                raw_json TEXT
            )
            """
        )
        con.execute(
            """
            CREATE TABLE IF NOT EXISTS experiment_runs (
                experiment_id TEXT PRIMARY KEY,
                run_at TEXT NOT NULL,
                parent_version TEXT,
                candidate_version TEXT,
                change_note TEXT,
                test_set_note TEXT,
                parent_prompt TEXT,
                candidate_prompt TEXT,
                model TEXT,
                target_words INTEGER,
                overlap_pct INTEGER,
                max_windows_per_file INTEGER,
                parent_files INTEGER,
                candidate_files INTEGER,
                parent_mean_ai REAL,
                candidate_mean_ai REAL,
                delta_ai REAL,
                candidate_worst_ai REAL,
                candidate_max_structure_similarity REAL,
                structure_similarity_limit REAL,
                verdict TEXT
            )
            """
        )
        con.commit()


def save_results(
    rows: list[dict[str, Any]],
    *,
    experiment_name: str,
    mode: str,
    model: str,
    experiment_id: str | None = None,
) -> str:
    init_db()
    experiment_id = experiment_id or str(uuid.uuid4())
    run_at = datetime.now(timezone.utc).isoformat()

    with sqlite3.connect(DB_PATH) as con:
        for r in rows:
            con.execute(
                """
                INSERT INTO scan_results (
                    experiment_id, experiment_name, run_at, mode, model,
                    source_name, expected_label, window_id, target_words, actual_words,
                    sentence_start, sentence_end, prediction, headline,
                    fraction_ai, fraction_ai_assisted, fraction_human,
                    mean_ai_involvement, max_humanizer_score, any_humanized,
                    window_confidence, version, dashboard_link, text, raw_json
                ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    experiment_id,
                    experiment_name,
                    run_at,
                    mode,
                    model,
                    r.get("source_name"),
                    r.get("expected_label"),
                    r.get("window_id"),
                    r.get("target_words"),
                    r.get("actual_words"),
                    r.get("sentence_start"),
                    r.get("sentence_end"),
                    r.get("prediction"),
                    r.get("headline"),
                    r.get("fraction_ai"),
                    r.get("fraction_ai_assisted"),
                    r.get("fraction_human"),
                    r.get("mean_ai_involvement"),
                    r.get("max_humanizer_score"),
                    1 if r.get("any_humanized") else 0,
                    r.get("window_confidence"),
                    r.get("version"),
                    r.get("dashboard_link"),
                    r.get("text"),
                    json.dumps(r.get("raw_result") or {}, ensure_ascii=False),
                ),
            )
        con.commit()
    return experiment_id


def save_experiment_run(record: dict[str, Any]) -> None:
    init_db()
    cols = [
        "experiment_id", "run_at", "parent_version", "candidate_version",
        "change_note", "test_set_note", "parent_prompt", "candidate_prompt",
        "model", "target_words", "overlap_pct", "max_windows_per_file",
        "parent_files", "candidate_files", "parent_mean_ai", "candidate_mean_ai",
        "delta_ai", "candidate_worst_ai", "candidate_max_structure_similarity",
        "structure_similarity_limit", "verdict"
    ]
    values = [record.get(c) for c in cols]
    placeholders = ",".join(["?"] * len(cols))
    with sqlite3.connect(DB_PATH) as con:
        con.execute(
            f"INSERT OR REPLACE INTO experiment_runs ({','.join(cols)}) VALUES ({placeholders})",
            values,
        )
        con.commit()


def load_history() -> pd.DataFrame:
    init_db()
    with sqlite3.connect(DB_PATH) as con:
        return pd.read_sql_query(
            "SELECT * FROM scan_results ORDER BY id DESC",
            con,
        )


def load_experiment_history() -> pd.DataFrame:
    init_db()
    with sqlite3.connect(DB_PATH) as con:
        return pd.read_sql_query(
            "SELECT * FROM experiment_runs ORDER BY run_at DESC",
            con,
        )


# -------------------------
# Analysis / display
# -------------------------


def results_dataframe(rows: list[dict[str, Any]]) -> pd.DataFrame:
    if not rows:
        return pd.DataFrame()
    display_rows = []
    for r in rows:
        d = {k: v for k, v in r.items() if k != "raw_result"}
        display_rows.append(d)
    return pd.DataFrame(display_rows)


def calibration_summary(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return df
    usable = df[df["expected_label"].isin(["Human", "AI"])].copy()
    if usable.empty:
        return pd.DataFrame()

    records = []
    for (label, size), group in usable.groupby(["expected_label", "target_words"]):
        n = len(group)
        records.append(
            {
                "Expected": label,
                "Target words": int(size),
                "N": n,
                "Human %": 100 * (group["prediction"] == "Human").mean(),
                "Mixed %": 100 * (group["prediction"] == "Mixed").mean(),
                "AI %": 100 * (group["prediction"] == "AI").mean(),
                "Mean human fraction": group["fraction_human"].mean(),
                "Mean AI fraction": group["fraction_ai"].mean(),
                "Mean AI involvement": group["mean_ai_involvement"].mean(),
            }
        )
    return pd.DataFrame(records).sort_values(["Target words", "Expected"])


def separation_summary(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame()
    records = []
    for size, group in df.groupby("target_words"):
        human = group[group["expected_label"] == "Human"]
        ai = group[group["expected_label"] == "AI"]
        if human.empty or ai.empty:
            continue
        human_correct = (human["prediction"] == "Human").mean()
        ai_correct = (ai["prediction"] == "AI").mean()
        records.append(
            {
                "Target words": int(size),
                "Human correctly Human %": 100 * human_correct,
                "AI correctly AI %": 100 * ai_correct,
                "Balanced decisive accuracy %": 100 * (human_correct + ai_correct) / 2,
                "Human N": len(human),
                "AI N": len(ai),
            }
        )
    return pd.DataFrame(records).sort_values("Target words")


def recommend_size(sep: pd.DataFrame, human_threshold: float, ai_threshold: float) -> str:
    if sep.empty:
        return "Not enough labeled Human and AI data to recommend a window size."
    candidates = sep[
        (sep["Human correctly Human %"] >= human_threshold * 100)
        & (sep["AI correctly AI %"] >= ai_threshold * 100)
    ]
    if candidates.empty:
        return (
            "No tested size met both thresholds. That is useful: either test larger windows, "
            "add more control texts, or accept a lower screening threshold for early experiments."
        )
    size = int(candidates.iloc[0]["Target words"])
    return (
        f"Smallest tested size meeting both thresholds: about **{size} words** "
        "(sentence-aligned, so individual windows may be somewhat longer)."
    )


def show_result_table(df: pd.DataFrame, key: str) -> None:
    if df.empty:
        st.info("No results yet.")
        return
    preferred = [
        "source_name",
        "expected_label",
        "target_words",
        "actual_words",
        "prediction",
        "fraction_human",
        "fraction_ai_assisted",
        "fraction_ai",
        "mean_ai_involvement",
        "max_humanizer_score",
        "window_confidence",
        "sentence_start",
        "sentence_end",
        "text",
    ]
    cols = [c for c in preferred if c in df.columns]
    st.dataframe(df[cols], use_container_width=True, hide_index=True, key=key)


def source_docs_from_uploads(files: Iterable[Any], expected_label: str) -> tuple[list[SourceDoc], list[str]]:
    docs = []
    errors = []
    for f in files or []:
        try:
            text = extract_text_from_upload(f)
            if count_words(text) < MIN_PANGRAM_WORDS:
                errors.append(f"{f.name}: fewer than {MIN_PANGRAM_WORDS} words after extraction")
                continue
            docs.append(SourceDoc(name=f.name, expected_label=expected_label, text=text))
        except Exception as exc:
            errors.append(f"{f.name}: {exc}")
    return docs, errors


def get_connected_client() -> tuple[Any | None, str | None]:
    return st.session_state.get("pangram_client"), st.session_state.get("pangram_model")


# -------------------------
# Experiment Lab analysis
# -------------------------


def weighted_mean(values: pd.Series, weights: pd.Series) -> float | None:
    usable = pd.DataFrame({"v": values, "w": weights}).dropna()
    if usable.empty or usable["w"].sum() <= 0:
        return None
    return float((usable["v"] * usable["w"]).sum() / usable["w"].sum())


def experiment_file_summary(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame()
    records = []
    for (side, source_name), group in df.groupby(["expected_label", "source_name"], sort=False):
        counts = group["prediction"].value_counts().to_dict()
        records.append(
            {
                "Version": side,
                "File": source_name,
                "Windows": len(group),
                "Submitted words": int(group["actual_words"].sum()),
                "Weighted AI fraction": weighted_mean(group["fraction_ai"], group["actual_words"]),
                "Weighted AI involvement": weighted_mean(group["mean_ai_involvement"], group["actual_words"]),
                "Human %": 100 * counts.get("Human", 0) / len(group),
                "Mixed %": 100 * counts.get("Mixed", 0) / len(group),
                "AI %": 100 * counts.get("AI", 0) / len(group),
            }
        )
    return pd.DataFrame(records)


def side_summary(file_summary: pd.DataFrame) -> pd.DataFrame:
    if file_summary.empty:
        return pd.DataFrame()
    records = []
    for version, group in file_summary.groupby("Version", sort=False):
        vals = group["Weighted AI fraction"].dropna()
        inv = group["Weighted AI involvement"].dropna()
        records.append(
            {
                "Version": version,
                "Files": len(group),
                "Mean AI fraction": float(vals.mean()) if not vals.empty else None,
                "Worst-file AI fraction": float(vals.max()) if not vals.empty else None,
                "Best-file AI fraction": float(vals.min()) if not vals.empty else None,
                "AI-fraction stdev": float(vals.std(ddof=0)) if len(vals) > 1 else 0.0,
                "Mean AI involvement": float(inv.mean()) if not inv.empty else None,
                "Human windows %": float(group["Human %"].mean()),
                "Mixed windows %": float(group["Mixed %"].mean()),
                "AI windows %": float(group["AI %"].mean()),
            }
        )
    return pd.DataFrame(records)


def _bucket_word_count(n: int) -> str:
    if n <= 25:
        return "A"
    if n <= 50:
        return "B"
    if n <= 90:
        return "C"
    if n <= 140:
        return "D"
    return "E"


def _bucket_sentence_count(n: int) -> str:
    return str(n) if n <= 4 else "5+"


def _bucket_sentence_words(n: int) -> str:
    if n <= 7:
        return "XS"
    if n <= 14:
        return "S"
    if n <= 24:
        return "M"
    if n <= 40:
        return "L"
    return "XL"


def structure_signature(text: str) -> dict[str, Any]:
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", clean_text(text)) if p.strip()]
    para_tokens: list[str] = []
    sentence_tokens: list[str] = []
    for p in paragraphs:
        sents = split_sentences(p)
        dialogue = p.lstrip().startswith(('"', '“', "'", '‘'))
        para_tokens.append(
            f"{_bucket_word_count(count_words(p))}|{_bucket_sentence_count(max(1, len(sents)))}|{'D' if dialogue else 'N'}"
        )
        for sent in sents:
            sentence_tokens.append(_bucket_sentence_words(count_words(sent)))
    return {
        "paragraphs": len(paragraphs),
        "sentences": len(sentence_tokens),
        "para_tokens": para_tokens,
        "sentence_tokens": sentence_tokens,
    }


def structure_similarity(text_a: str, text_b: str) -> dict[str, float]:
    a = structure_signature(text_a)
    b = structure_signature(text_b)
    para = difflib.SequenceMatcher(None, a["para_tokens"], b["para_tokens"], autojunk=False).ratio()
    sent = difflib.SequenceMatcher(None, a["sentence_tokens"], b["sentence_tokens"], autojunk=False).ratio()
    combined = 0.70 * para + 0.30 * sent
    return {
        "paragraph_similarity": float(para),
        "sentence_shape_similarity": float(sent),
        "combined_similarity": float(combined),
    }


def candidate_diversity_table(docs: list[SourceDoc]) -> pd.DataFrame:
    if len(docs) < 2:
        return pd.DataFrame()
    rows = []
    for i in range(len(docs)):
        for j in range(i + 1, len(docs)):
            sim = structure_similarity(docs[i].text, docs[j].text)
            rows.append(
                {
                    "File A": docs[i].name,
                    "File B": docs[j].name,
                    "Paragraph-shape similarity": sim["paragraph_similarity"],
                    "Sentence-shape similarity": sim["sentence_shape_similarity"],
                    "Combined structural similarity": sim["combined_similarity"],
                }
            )
    return pd.DataFrame(rows).sort_values("Combined structural similarity", ascending=False)


def build_experiment_windows(
    docs: list[SourceDoc],
    target_words: int,
    overlap_pct: int,
    cap_per_file: int,
) -> list[TextWindow]:
    out: list[TextWindow] = []
    for doc in docs:
        out.extend(
            evenly_cap(
                build_sentence_windows(doc, target_words, overlap_pct / 100.0),
                cap_per_file,
            )
        )
    return out


def prompt_diff(parent_prompt: str, candidate_prompt: str) -> str:
    if not parent_prompt.strip() and not candidate_prompt.strip():
        return ""
    diff = difflib.unified_diff(
        parent_prompt.splitlines(),
        candidate_prompt.splitlines(),
        fromfile="parent prompt",
        tofile="candidate prompt",
        lineterm="",
    )
    return "\n".join(diff)


def experiment_verdict(delta_ai: float | None) -> str:
    if delta_ai is None:
        return "NO COMPARISON"
    if delta_ai <= -0.10:
        return "STRONG IMPROVEMENT"
    if delta_ai <= -0.05:
        return "PROMISING"
    if delta_ai < 0.05:
        return "INCONCLUSIVE"
    return "WORSE"


def compact_handoff(
    parent_version: str,
    candidate_version: str,
    change_note: str,
    file_summary: pd.DataFrame,
    side_summary_df: pd.DataFrame,
    diversity_df: pd.DataFrame,
    structure_limit: float,
) -> str:
    lines = [
        f"Pangram Experiment Lab: {parent_version} → {candidate_version}",
        f"Change tested: {change_note or '(not entered)'}",
        "",
        "Version summary:",
    ]
    if not side_summary_df.empty:
        for _, r in side_summary_df.iterrows():
            lines.append(
                f"- {r['Version']}: {int(r['Files'])} file(s), mean AI fraction {100*r['Mean AI fraction']:.1f}%, "
                f"worst file {100*r['Worst-file AI fraction']:.1f}%"
            )
    if not file_summary.empty:
        lines.append("")
        lines.append("Files:")
        for _, r in file_summary.iterrows():
            lines.append(
                f"- {r['Version']} / {r['File']}: {r['Windows']} windows, weighted AI {100*r['Weighted AI fraction']:.1f}%"
            )
    if not diversity_df.empty:
        max_sim = float(diversity_df["Combined structural similarity"].max())
        lines += [
            "",
            f"Candidate max pairwise structural similarity: {100*max_sim:.1f}% "
            f"(experimental warning line {100*structure_limit:.0f}%).",
        ]
    else:
        lines += ["", "Candidate structural diversity: not testable from fewer than 2 candidate files."]
    return "\n".join(lines)


# -------------------------
# Streamlit app
# -------------------------

st.set_page_config(page_title=APP_TITLE, layout="wide")
init_db()

st.title(f"{APP_TITLE} {APP_VERSION}")
st.caption(
    "A prompt-development lab that uses Pangram as a measurement instrument while separately checking "
    "whether candidate outputs are converging on the same structural skeleton."
)

with st.sidebar:
    st.header("Pangram connection")
    api_key = get_secret_api_key()

    if Pangram is None:
        st.error(
            "pangram-sdk is not installed. Make sure requirements.txt is in the GitHub repo, "
            "then reboot the Streamlit app."
        )
    elif not api_key:
        st.error("PANGRAM_API_KEY is not set in Streamlit Secrets.")
        st.caption("In Streamlit Cloud: Manage app → Settings → Secrets, then add:")
        st.code('PANGRAM_API_KEY = "paste-your-key-here"', language="toml")
        st.caption("Save the secret and let Streamlit rerun the app. Do not put the key in GitHub.")
    else:
        # Auto-connect once per Streamlit session. No command line and no key-pasting in the app.
        if st.session_state.get("pangram_client") is None:
            try:
                with st.spinner("Connecting to Pangram…"):
                    client, models = connect_pangram(api_key)
                st.session_state["pangram_client"] = client
                st.session_state["pangram_models"] = models
                if "pangram-4" in models:
                    st.session_state["pangram_model"] = "pangram-4"
                elif models:
                    st.session_state["pangram_model"] = models[0]
            except Exception as exc:
                st.session_state.pop("pangram_client", None)
                st.session_state.pop("pangram_models", None)
                st.session_state.pop("pangram_model", None)
                st.error(f"Pangram connection failed: {exc}")

        if st.session_state.get("pangram_client") is not None:
            st.success("API key loaded from Streamlit Secrets.")

            if st.button("Refresh Pangram models", use_container_width=True):
                try:
                    with st.spinner("Refreshing Pangram models…"):
                        client, models = connect_pangram(api_key)
                    st.session_state["pangram_client"] = client
                    st.session_state["pangram_models"] = models
                    current = st.session_state.get("pangram_model")
                    if current not in models:
                        st.session_state["pangram_model"] = (
                            "pangram-4" if "pangram-4" in models else (models[0] if models else None)
                        )
                    st.rerun()
                except Exception as exc:
                    st.error(f"Could not refresh models: {exc}")

    models = st.session_state.get("pangram_models", [])
    if models:
        current_model = st.session_state.get("pangram_model")
        idx = models.index(current_model) if current_model in models else 0
        selected = st.selectbox("Model", models, index=idx)
        st.session_state["pangram_model"] = selected
        st.caption("This list is read from the models currently enabled for your Pangram API key.")
    elif api_key and Pangram is not None:
        st.caption("No Pangram model list is available yet.")

    st.divider()
    st.caption(
        "Pangram 4 accepts prose samples of at least 50 words. This app enforces that minimum and "
        "builds windows on sentence boundaries."
    )

experiment_tab, cal_tab, quick_tab, ab_tab, history_tab = st.tabs(
    [
        "1 · Experiment Lab",
        "2 · Corpus calibration",
        "3 · 150-word microscope",
        "4 · Legacy A/B",
        "5 · History",
    ]
)


# -------------------------
# Tab 1: Experiment Lab
# -------------------------
with experiment_tab:
    st.subheader("Prompt experiment + generalization check")
    st.write(
        "The detector score is only one objective. A candidate should improve Pangram across more than one "
        "sample without making different chapters collapse onto the same paragraph/sentence skeleton."
    )
    st.info(
        "Development test: one parent file + one candidate file is enough to compare 6D → 6E. "
        "Promotion test: use 3 or more candidate outputs from different chapters and different donor structures."
    )

    v1, v2 = st.columns(2)
    with v1:
        parent_version = st.text_input("Parent version", value="6D", key="lab_parent_version")
    with v2:
        candidate_version = st.text_input("Candidate version", value="6E", key="lab_candidate_version")

    change_note = st.text_area(
        "What changed in the candidate prompt?",
        placeholder="One controlled change only, if possible.",
        height=90,
        key="lab_change_note",
    )
    test_set_note = st.text_input(
        "Test-set note",
        placeholder="Example: Ch. 3 + three different donors; holdout donors not used to design 6E",
        key="lab_test_set_note",
    )

    with st.expander("Store the two prompts with this experiment", expanded=False):
        p1, p2 = st.columns(2)
        with p1:
            parent_prompt = st.text_area("Parent prompt", height=300, key="lab_parent_prompt")
        with p2:
            candidate_prompt = st.text_area("Candidate prompt", height=300, key="lab_candidate_prompt")
        diff_text = prompt_diff(parent_prompt, candidate_prompt)
        if diff_text:
            st.markdown("**Prompt diff**")
            st.code(diff_text, language="diff")

    u1, u2 = st.columns(2)
    with u1:
        parent_uploads = st.file_uploader(
            f"{parent_version} output(s)",
            type=["docx", "txt", "md"],
            accept_multiple_files=True,
            key="lab_parent_uploads",
            help="For the first 6D→6E test, one 6D chapter is enough. Later, use matched sets when possible.",
        )
    with u2:
        candidate_uploads = st.file_uploader(
            f"{candidate_version} output(s)",
            type=["docx", "txt", "md"],
            accept_multiple_files=True,
            key="lab_candidate_uploads",
            help="For promotion/generalization, upload at least 3 outputs created from different donor structures.",
        )

    parent_docs, parent_errors = source_docs_from_uploads(parent_uploads, parent_version)
    candidate_docs, candidate_errors = source_docs_from_uploads(candidate_uploads, candidate_version)
    for err in parent_errors + candidate_errors:
        st.warning(err)

    s1, s2, s3, s4 = st.columns(4)
    with s1:
        lab_size = st.selectbox(
            "Pangram window",
            ALL_SAMPLE_SIZES,
            index=ALL_SAMPLE_SIZES.index(EXPERIMENT_WINDOW_WORDS),
            key="lab_size",
        )
    with s2:
        lab_overlap = st.slider(
            "Overlap",
            0,
            75,
            EXPERIMENT_OVERLAP_PCT,
            25,
            key="lab_overlap",
        )
    with s3:
        lab_cap = st.number_input(
            "Max windows / file",
            min_value=1,
            max_value=100,
            value=EXPERIMENT_MAX_WINDOWS_PER_FILE,
            step=1,
            key="lab_cap",
        )
    with s4:
        structure_limit = st.slider(
            "Structure warning line",
            0.60,
            0.95,
            DEFAULT_STRUCTURE_SIMILARITY_LIMIT,
            0.01,
            key="lab_structure_limit",
            help="Experimental heuristic, not a literary-quality score. Higher means two outputs have more similar paragraph/sentence shape sequences.",
        )

    parent_windows = build_experiment_windows(parent_docs, int(lab_size), int(lab_overlap), int(lab_cap))
    candidate_windows = build_experiment_windows(candidate_docs, int(lab_size), int(lab_overlap), int(lab_cap))
    lab_windows = parent_windows + candidate_windows

    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Parent files", len(parent_docs))
    m2.metric("Candidate files", len(candidate_docs))
    m3.metric("Pangram windows", len(lab_windows))
    m4.metric("Submitted words", f"{sum(w.actual_words for w in lab_windows):,}")
    show_cost_estimate(lab_windows, key="lab_cost")

    pre_diversity = candidate_diversity_table(candidate_docs)
    if len(candidate_docs) >= 2:
        with st.expander("Pre-scan structural diversity check", expanded=False):
            st.dataframe(
                pre_diversity.style.format({
                    "Paragraph-shape similarity": "{:.1%}",
                    "Sentence-shape similarity": "{:.1%}",
                    "Combined structural similarity": "{:.1%}",
                }),
                use_container_width=True,
                hide_index=True,
            )
            max_pre = float(pre_diversity["Combined structural similarity"].max())
            if max_pre >= structure_limit:
                st.warning(
                    f"At least one candidate pair is {max_pre:.1%} structurally similar. "
                    "That does not prove bad prose, but it is exactly the kind of repeated skeleton this project needs to catch."
                )
            else:
                st.success(f"No candidate pair crosses the current {structure_limit:.0%} structure-warning line.")
    elif len(candidate_docs) == 1:
        st.caption("One candidate file can test Pangram improvement, but it cannot test cross-chapter structural repetition.")

    client, model = get_connected_client()
    run_disabled = not (client and model and parent_windows and candidate_windows)
    if st.button(
        f"Run {parent_version} → {candidate_version} experiment",
        type="primary",
        disabled=run_disabled,
        key="lab_run",
    ):
        try:
            successes, failures = run_bulk_scan(client, model, lab_windows)
            run_id = str(uuid.uuid4())
            save_results(
                successes,
                experiment_name=f"{parent_version} → {candidate_version}",
                mode="Experiment Lab",
                model=model,
                experiment_id=run_id,
            )
            df = results_dataframe(successes)
            file_sum = experiment_file_summary(df)
            side_sum = side_summary(file_sum)
            parent_row = side_sum[side_sum["Version"] == parent_version]
            candidate_row = side_sum[side_sum["Version"] == candidate_version]
            parent_mean = float(parent_row.iloc[0]["Mean AI fraction"]) if not parent_row.empty else None
            candidate_mean = float(candidate_row.iloc[0]["Mean AI fraction"]) if not candidate_row.empty else None
            delta_ai = candidate_mean - parent_mean if parent_mean is not None and candidate_mean is not None else None
            candidate_worst = float(candidate_row.iloc[0]["Worst-file AI fraction"]) if not candidate_row.empty else None
            diversity_df = candidate_diversity_table(candidate_docs)
            max_sim = float(diversity_df["Combined structural similarity"].max()) if not diversity_df.empty else None
            verdict = experiment_verdict(delta_ai)
            save_experiment_run(
                {
                    "experiment_id": run_id,
                    "run_at": datetime.now(timezone.utc).isoformat(),
                    "parent_version": parent_version,
                    "candidate_version": candidate_version,
                    "change_note": change_note,
                    "test_set_note": test_set_note,
                    "parent_prompt": parent_prompt,
                    "candidate_prompt": candidate_prompt,
                    "model": model,
                    "target_words": int(lab_size),
                    "overlap_pct": int(lab_overlap),
                    "max_windows_per_file": int(lab_cap),
                    "parent_files": len(parent_docs),
                    "candidate_files": len(candidate_docs),
                    "parent_mean_ai": parent_mean,
                    "candidate_mean_ai": candidate_mean,
                    "delta_ai": delta_ai,
                    "candidate_worst_ai": candidate_worst,
                    "candidate_max_structure_similarity": max_sim,
                    "structure_similarity_limit": float(structure_limit),
                    "verdict": verdict,
                }
            )
            st.session_state["lab_last"] = {
                "run_id": run_id,
                "rows": successes,
                "failures": failures,
                "file_summary": file_sum,
                "side_summary": side_sum,
                "diversity": diversity_df,
                "parent_version": parent_version,
                "candidate_version": candidate_version,
                "change_note": change_note,
                "structure_limit": float(structure_limit),
                "verdict": verdict,
                "delta_ai": delta_ai,
            }
        except Exception as exc:
            st.error(f"Experiment failed: {exc}")

    lab_last = st.session_state.get("lab_last")
    if lab_last:
        st.divider()
        st.subheader("Experiment result")
        side_sum = lab_last["side_summary"]
        file_sum = lab_last["file_summary"]
        diversity_df = lab_last["diversity"]
        delta_ai = lab_last["delta_ai"]
        verdict = lab_last["verdict"]

        if delta_ai is not None:
            r1, r2, r3 = st.columns(3)
            p_row = side_sum[side_sum["Version"] == lab_last["parent_version"]].iloc[0]
            c_row = side_sum[side_sum["Version"] == lab_last["candidate_version"]].iloc[0]
            r1.metric(f"{lab_last['parent_version']} mean AI", f"{100*p_row['Mean AI fraction']:.1f}%")
            r2.metric(f"{lab_last['candidate_version']} mean AI", f"{100*c_row['Mean AI fraction']:.1f}%")
            r3.metric("Candidate change", f"{100*delta_ai:+.1f} points", delta_color="inverse")
            if verdict in {"STRONG IMPROVEMENT", "PROMISING"}:
                st.success(f"Detector result: **{verdict}**")
            elif verdict == "INCONCLUSIVE":
                st.info("Detector result: **INCONCLUSIVE**")
            else:
                st.error(f"Detector result: **{verdict}**")

        st.markdown("**File-level Pangram results**")
        st.dataframe(
            file_sum.style.format({
                "Weighted AI fraction": "{:.1%}",
                "Weighted AI involvement": "{:.1%}",
                "Human %": "{:.1f}",
                "Mixed %": "{:.1f}",
                "AI %": "{:.1f}",
            }),
            use_container_width=True,
            hide_index=True,
        )

        st.markdown("**Generalization / diversity gate**")
        candidate_count = int((file_sum["Version"] == lab_last["candidate_version"]).sum())
        if candidate_count < 2:
            st.warning("Not tested: one candidate output cannot tell us whether chapters/books are converging on one skeleton.")
        else:
            max_sim = float(diversity_df["Combined structural similarity"].max())
            st.dataframe(
                diversity_df.style.format({
                    "Paragraph-shape similarity": "{:.1%}",
                    "Sentence-shape similarity": "{:.1%}",
                    "Combined structural similarity": "{:.1%}",
                }),
                use_container_width=True,
                hide_index=True,
            )
            if max_sim >= lab_last["structure_limit"]:
                st.error(
                    f"Structure warning: max pairwise similarity is {max_sim:.1%}. "
                    "Do not promote the prompt on Pangram score alone."
                )
            elif candidate_count < 3:
                st.info(
                    f"No pair crossed the warning line, but use at least 3 different chapter/donor outputs before promotion. "
                    f"Current max similarity: {max_sim:.1%}."
                )
            else:
                st.success(
                    f"Diversity check passed at the current experimental threshold. Max candidate-pair similarity: {max_sim:.1%}."
                )

        with st.expander("Every Pangram window"):
            show_result_table(results_dataframe(lab_last["rows"]), "lab_results_df")

        handoff = compact_handoff(
            lab_last["parent_version"],
            lab_last["candidate_version"],
            lab_last["change_note"],
            file_sum,
            side_sum,
            diversity_df,
            lab_last["structure_limit"],
        )
        st.text_area("Copy back to ChatGPT", value=handoff, height=240, key="lab_handoff")
        st.download_button(
            "Download experiment CSV",
            results_dataframe(lab_last["rows"]).to_csv(index=False).encode("utf-8"),
            file_name=f"pangram_experiment_{lab_last['parent_version']}_to_{lab_last['candidate_version']}.csv",
            mime="text/csv",
        )
        if lab_last["failures"]:
            st.warning(f"{len(lab_last['failures'])} Pangram item(s) failed.")
            st.dataframe(pd.DataFrame(lab_last["failures"]), use_container_width=True, hide_index=True)


# -------------------------
# Tab 2: Calibration
# -------------------------
with cal_tab:
    st.subheader("Find the smallest useful test window")
    st.write(
        "Load writing you know is human and writing you know was generated by AI. The app creates "
        "mechanical, sentence-aligned windows at several sizes and sends them to Pangram in bulk."
    )

    c1, c2 = st.columns(2)
    with c1:
        human_files = st.file_uploader(
            "Known-human corpus",
            type=["docx", "txt", "md"],
            accept_multiple_files=True,
            key="human_uploads",
        )
    with c2:
        ai_files = st.file_uploader(
            "Known-AI corpus",
            type=["docx", "txt", "md"],
            accept_multiple_files=True,
            key="ai_uploads",
        )

    human_docs, human_errors = source_docs_from_uploads(human_files, "Human")
    ai_docs, ai_errors = source_docs_from_uploads(ai_files, "AI")
    for err in human_errors + ai_errors:
        st.warning(err)

    settings1, settings2, settings3 = st.columns(3)
    with settings1:
        sample_sizes = st.multiselect(
            "Target window sizes (words)",
            ALL_SAMPLE_SIZES,
            default=DEFAULT_SAMPLE_SIZES,
        )
    with settings2:
        overlap_pct = st.slider("Window overlap", 0, 75, DEFAULT_OVERLAP_PCT, 25)
    with settings3:
        cap_per = st.number_input(
            "Max windows / file / size",
            min_value=1,
            max_value=200,
            value=CALIBRATION_DEFAULT_MAX_WINDOWS,
            step=1,
            help="Windows are selected evenly across each file when this cap is reached.",
        )

    docs = human_docs + ai_docs
    windows = make_calibration_windows(
        docs,
        [s for s in sample_sizes if s >= MIN_PANGRAM_WORDS],
        overlap_pct / 100.0,
        int(cap_per),
    ) if docs and sample_sizes else []

    total_words = sum(w.actual_words for w in windows)
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Human files", len(human_docs))
    m2.metric("AI files", len(ai_docs))
    m3.metric("Pangram samples", len(windows))
    m4.metric("Words submitted", f"{total_words:,}")

    if windows:
        preview = pd.DataFrame([asdict(w) for w in windows])
        with st.expander("Preview the mechanical sample plan"):
            st.dataframe(
                preview[["source_name", "expected_label", "target_words", "actual_words", "sentence_start", "sentence_end"]],
                use_container_width=True,
                hide_index=True,
            )
        show_cost_estimate(windows, key="cal_cost")

    experiment_name = st.text_input(
        "Experiment name",
        value=f"Calibration {datetime.now().strftime('%Y-%m-%d')}",
        key="cal_experiment_name",
    )

    client, model = get_connected_client()
    run_disabled = not (client and model and windows and human_docs and ai_docs)
    if st.button("Run calibration through Pangram", type="primary", disabled=run_disabled):
        try:
            successes, failures = run_bulk_scan(client, model, windows)
            exp_id = save_results(
                successes,
                experiment_name=experiment_name,
                mode="Calibration",
                model=model,
            )
            st.session_state["cal_last_rows"] = successes
            st.session_state["cal_last_failures"] = failures
            st.session_state["cal_last_experiment_id"] = exp_id
        except Exception as exc:
            st.error(f"Pangram run failed: {exc}")

    rows = st.session_state.get("cal_last_rows", [])
    failures = st.session_state.get("cal_last_failures", [])
    if rows:
        df = results_dataframe(rows)
        st.divider()
        st.subheader("Calibration result")

        summary = calibration_summary(df)
        sep = separation_summary(df)

        if not sep.empty:
            st.markdown("**Separation by test length**")
            st.dataframe(sep, use_container_width=True, hide_index=True)
            chart = sep.set_index("Target words")[["Human correctly Human %", "AI correctly AI %"]]
            st.line_chart(chart)

            tc1, tc2 = st.columns(2)
            with tc1:
                human_thresh = st.slider("Required Human-control success", 0.50, 1.00, 0.90, 0.05)
            with tc2:
                ai_thresh = st.slider("Required AI-control success", 0.50, 1.00, 0.90, 0.05)
            st.info(recommend_size(sep, human_thresh, ai_thresh))

        with st.expander("Detailed calibration summary"):
            st.dataframe(summary, use_container_width=True, hide_index=True)

        with st.expander("Every tested window"):
            show_result_table(df, "cal_results_df")

        st.download_button(
            "Download calibration CSV",
            data=df.to_csv(index=False).encode("utf-8"),
            file_name="pangram_calibration_results.csv",
            mime="text/csv",
        )

        if failures:
            st.warning(f"{len(failures)} Pangram item(s) failed.")
            st.dataframe(pd.DataFrame(failures), use_container_width=True, hide_index=True)

    elif not run_disabled and not rows:
        st.caption("Ready. Run the calibration when you want to spend the API calls.")
    elif not client:
        st.info("Connect to Pangram in the sidebar first.")
    elif not (human_docs and ai_docs):
        st.info("Add at least one known-human and one known-AI document.")


# -------------------------
# Tab 3: 150-word microscope
# -------------------------
with quick_tab:
    st.subheader("150-word microscope")
    st.write(
        "Use this for local diagnosis, not for ranking prompt versions. The 6C/6D calibration showed that "
        "150-word windows can disagree with the whole-document direction."
    )

    quick_upload = st.file_uploader(
        "Optional DOCX / TXT / MD",
        type=["docx", "txt", "md"],
        accept_multiple_files=False,
        key="quick_upload",
    )
    default_quick_text = ""
    quick_name = "pasted_text.txt"
    if quick_upload is not None:
        try:
            default_quick_text = extract_text_from_upload(quick_upload)
            quick_name = quick_upload.name
        except Exception as exc:
            st.error(str(exc))

    quick_text = st.text_area(
        "Text to test",
        value=default_quick_text,
        height=260,
        key="quick_text",
    )

    q1, q2, q3 = st.columns(3)
    with q1:
        q_size = st.selectbox("Target words", ALL_SAMPLE_SIZES, index=ALL_SAMPLE_SIZES.index(CALIBRATED_WINDOW_WORDS))
    with q2:
        q_overlap = st.slider("Overlap", 0, 75, DEFAULT_OVERLAP_PCT, 25, key="quick_overlap")
    with q3:
        q_cap = st.number_input("Max windows", 1, 100, QUICK_DEFAULT_MAX_WINDOWS, 1, key="quick_cap")

    qdoc = SourceDoc(name=quick_name, expected_label="Unknown", text=clean_text(quick_text))
    qwindows = evenly_cap(
        build_sentence_windows(qdoc, int(q_size), q_overlap / 100.0),
        int(q_cap),
    ) if count_words(qdoc.text) >= MIN_PANGRAM_WORDS else []

    st.caption(
        f"{count_words(qdoc.text):,} source words → {len(qwindows)} Pangram window(s) → "
        f"{sum(w.actual_words for w in qwindows):,} submitted words."
    )
    show_cost_estimate(qwindows, key="quick_cost")

    qexp = st.text_input("Experiment name", value="Quick scan", key="quick_exp")
    client, model = get_connected_client()
    if st.button(
        "Scan these windows",
        type="primary",
        disabled=not (client and model and qwindows),
        key="quick_run",
    ):
        try:
            successes, failures = run_bulk_scan(client, model, qwindows)
            save_results(successes, experiment_name=qexp, mode="Quick", model=model)
            st.session_state["quick_last_rows"] = successes
            st.session_state["quick_last_failures"] = failures
        except Exception as exc:
            st.error(f"Pangram run failed: {exc}")

    qrows = st.session_state.get("quick_last_rows", [])
    if qrows:
        qdf = results_dataframe(qrows)
        qcounts = qdf["prediction"].value_counts().to_dict()
        cols = st.columns(3)
        cols[0].metric("Human windows", qcounts.get("Human", 0))
        cols[1].metric("Mixed windows", qcounts.get("Mixed", 0))
        cols[2].metric("AI windows", qcounts.get("AI", 0))
        show_result_table(qdf, "quick_results_df")
        st.download_button(
            "Download quick-scan CSV",
            qdf.to_csv(index=False).encode("utf-8"),
            file_name="pangram_quick_scan.csv",
            mime="text/csv",
        )


# -------------------------
# Tab 4: Legacy A/B prompt test
# -------------------------
with ab_tab:
    st.subheader("Legacy A/B prompt experiment")
    st.write(
        "Paste comparable output from the current prompt and one candidate prompt. The app uses the same "
        "window settings on both and compares their Pangram distributions."
    )

    a_col, b_col = st.columns(2)
    with a_col:
        control_text = st.text_area("Control output", height=300, key="control_text")
    with b_col:
        candidate_text = st.text_area("Candidate output", height=300, key="candidate_text")

    ab1, ab2, ab3 = st.columns(3)
    with ab1:
        ab_size = st.selectbox("Target words", ALL_SAMPLE_SIZES, index=ALL_SAMPLE_SIZES.index(CALIBRATED_WINDOW_WORDS), key="ab_size")
    with ab2:
        ab_overlap = st.slider("Overlap", 0, 75, DEFAULT_OVERLAP_PCT, 25, key="ab_overlap")
    with ab3:
        ab_cap = st.number_input("Max windows per side", 1, 100, QUICK_DEFAULT_MAX_WINDOWS, 1, key="ab_cap")

    cdoc = SourceDoc("Control", "Control", clean_text(control_text))
    ndoc = SourceDoc("Candidate", "Candidate", clean_text(candidate_text))
    cwindows = evenly_cap(build_sentence_windows(cdoc, int(ab_size), ab_overlap / 100.0), int(ab_cap)) if count_words(cdoc.text) >= MIN_PANGRAM_WORDS else []
    nwindows = evenly_cap(build_sentence_windows(ndoc, int(ab_size), ab_overlap / 100.0), int(ab_cap)) if count_words(ndoc.text) >= MIN_PANGRAM_WORDS else []
    abwindows = cwindows + nwindows
    show_cost_estimate(abwindows, key="ab_cost")

    abexp = st.text_input("Experiment name", value="Prompt A/B", key="ab_exp")
    client, model = get_connected_client()
    if st.button(
        "Run A/B Pangram test",
        type="primary",
        disabled=not (client and model and cwindows and nwindows),
        key="ab_run",
    ):
        try:
            successes, failures = run_bulk_scan(client, model, abwindows)
            save_results(successes, experiment_name=abexp, mode="A/B", model=model)
            st.session_state["ab_last_rows"] = successes
            st.session_state["ab_last_failures"] = failures
        except Exception as exc:
            st.error(f"Pangram run failed: {exc}")

    abrows = st.session_state.get("ab_last_rows", [])
    if abrows:
        abdf = results_dataframe(abrows)
        compare = []
        for label, group in abdf.groupby("expected_label"):
            compare.append(
                {
                    "Version": label,
                    "Windows": len(group),
                    "Human %": 100 * (group["prediction"] == "Human").mean(),
                    "Mixed %": 100 * (group["prediction"] == "Mixed").mean(),
                    "AI %": 100 * (group["prediction"] == "AI").mean(),
                    "Mean human fraction": group["fraction_human"].mean(),
                    "Mean AI fraction": group["fraction_ai"].mean(),
                    "Mean AI involvement": group["mean_ai_involvement"].mean(),
                }
            )
        cmp_df = pd.DataFrame(compare)
        st.dataframe(cmp_df, use_container_width=True, hide_index=True)
        if not cmp_df.empty:
            st.bar_chart(cmp_df.set_index("Version")[["Human %", "Mixed %", "AI %"]])
        with st.expander("Every A/B window"):
            show_result_table(abdf, "ab_results_df")


# -------------------------
# Tab 5: History
# -------------------------
with history_tab:
    st.subheader("Experiment history")
    exp_hist = load_experiment_history()
    if not exp_hist.empty:
        st.markdown("**Experiment Lab runs**")
        exp_show = exp_hist.copy()
        for c in ["parent_mean_ai", "candidate_mean_ai", "delta_ai", "candidate_worst_ai", "candidate_max_structure_similarity"]:
            if c in exp_show.columns:
                exp_show[c] = pd.to_numeric(exp_show[c], errors="coerce")
        st.dataframe(exp_show, use_container_width=True, hide_index=True)
        st.download_button(
            "Download Experiment Lab history CSV",
            exp_show.to_csv(index=False).encode("utf-8"),
            file_name="pangram_experiment_lab_history.csv",
            mime="text/csv",
            key="download_lab_history",
        )
        st.divider()
    st.caption(
        "Streamlit Cloud can recycle an app's local filesystem. Treat this built-in SQLite history as "
        "working history, not permanent storage; download CSVs you want to keep."
    )
    hist = load_history()
    if hist.empty:
        st.info("No saved Pangram runs yet.")
    else:
        e1, e2, e3 = st.columns(3)
        e1.metric("Saved windows", len(hist))
        e2.metric("Experiments", hist["experiment_id"].nunique())
        e3.metric("Latest run", str(hist.iloc[0]["run_at"])[:19].replace("T", " "))

        names = ["All"] + sorted([x for x in hist["experiment_name"].dropna().unique()])
        selected_name = st.selectbox("Experiment", names)
        shown = hist if selected_name == "All" else hist[hist["experiment_name"] == selected_name]

        display_cols = [
            "run_at",
            "experiment_name",
            "mode",
            "model",
            "source_name",
            "expected_label",
            "target_words",
            "actual_words",
            "prediction",
            "fraction_human",
            "fraction_ai_assisted",
            "fraction_ai",
            "mean_ai_involvement",
            "max_humanizer_score",
            "text",
        ]
        st.dataframe(
            shown[[c for c in display_cols if c in shown.columns]],
            use_container_width=True,
            hide_index=True,
        )
        st.download_button(
            "Download history CSV",
            shown.to_csv(index=False).encode("utf-8"),
            file_name="pangram_microscope_history.csv",
            mime="text/csv",
        )

st.divider()
st.caption(
    "Interpretation rule: this is a measurement lab, not a rewriting loop. Keep the drafting model blind to Pangram results. "
    "Optimize for generalization across different chapters/donors, not for one detector-perfect structural template."
)
